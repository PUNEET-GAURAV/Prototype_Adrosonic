from pptx import Presentation
from pptx.util import Inches, Pt
import os

def create_ppt():
    prs = Presentation()
    
    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Verity-RAG Prototype Demo"
    subtitle.text = "Vector Database Design for Large-Scale Precision Retrieval\nAdrosonic Build\n"

    # Slide 2: Prototype Execution
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Prototype Execution Overview"
    tf = body_shape.text_frame
    tf.text = "1. Data Ingestion: Successfully ingested 10,000 MS MARCO passages."
    p = tf.add_paragraph()
    p.text = "2. IR Metrics Benchmarking: Completed Recall@5 and Hit@5 evaluation on 50 samples for dense and hybrid modes."
    p = tf.add_paragraph()
    p.text = "3. RAGAS Evaluation: Measured Context Precision & Context Recall using Groq LLM (openai/gpt-oss-20b)."
    p = tf.add_paragraph()
    p.text = "4. Streamlit Demo: Functional UI built to compare dense vs hybrid side-by-side."

    # Slide 3: Next Steps
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Gate 3 & Next Steps"
    tf = body_shape.text_frame
    tf.text = "1. Consolidate results into 'results/' directory (hardware, IR metrics, RAGAS, latency)."
    p = tf.add_paragraph()
    p.text = "2. Run 'make report' for final REPORT.md."
    p = tf.add_paragraph()
    p.text = "3. Code freeze and prepare for Judge Check workflow."

    # Save
    out_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "demo_presentation.pptx")
    prs.save(out_path)
    print(f"Presentation saved to {out_path}")

if __name__ == "__main__":
    create_ppt()
